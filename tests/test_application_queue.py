from datetime import date

from openpyxl import load_workbook

from src.application_queue import create_application_review_page
from src.reports.export_excel import export_daily_new_leads


def test_review_page_and_excel_link_are_created(tmp_path) -> None:
    jobs = [
        {
            "title": "Business Analyst",
            "company": "Example",
            "location": "Hyderabad",
            "job_url": "https://example.com/apply",
            "posted_days": "1",
            "application_status": "Answers needed",
            "unknown_questions": ["Years of experience"],
        }
    ]

    review_page = create_application_review_page(
        jobs,
        tmp_path,
        run_date=date(2026, 9, 17),
    )
    jobs[0]["application_progress"] = (
        f"{review_page.name}#1"
    )
    output = export_daily_new_leads(
        jobs,
        tmp_path,
        run_date=date(2026, 9, 17),
        freshness_windows=(7,),
    )[0]

    assert "Years of experience" in review_page.read_text()
    worksheet = load_workbook(output).active
    headers = [cell.value for cell in worksheet[1]]
    progress_column = headers.index("application_progress") + 1
    cell = worksheet.cell(row=2, column=progress_column)
    assert cell.hyperlink is not None
    assert cell.hyperlink.target.endswith("#1")
