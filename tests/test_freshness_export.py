from datetime import date

from openpyxl import load_workbook

from src.freshness import is_within_freshness, posted_age_days
from src.reports.export_excel import export_daily_new_leads


def test_posted_age_parses_relative_values() -> None:
    assert posted_age_days({"posted_date": "3 days ago"}) == 3
    assert posted_age_days({"posted_date": "2 weeks ago"}) == 14
    assert posted_age_days({"posted_date": "today"}) == 0
    assert posted_age_days({"posted_date": "unknown"}) is None


def test_daily_exports_separate_freshness_windows(tmp_path) -> None:
    jobs = [
        {"title": "Recent", "posted_days": "3", "job_url": "1"},
        {"title": "Monthly", "posted_days": "10", "job_url": "2"},
        {"title": "Stale", "posted_days": "31", "job_url": "3"},
        {"title": "Unknown", "posted_days": "", "job_url": "4"},
    ]

    assert is_within_freshness(jobs[0], 7)
    assert not is_within_freshness(jobs[2], 30)
    assert not is_within_freshness(jobs[3], 30)

    paths = export_daily_new_leads(
        jobs,
        tmp_path,
        run_date=date(2026, 9, 17),
        freshness_windows=(7, 30),
    )

    seven_day = load_workbook(paths[0]).active
    thirty_day = load_workbook(paths[1]).active

    assert seven_day.max_row == 2
    assert thirty_day.max_row == 3
    assert seven_day[2][0].value == "Recent"
    assert thirty_day[3][0].value == "Monthly"


def test_age_keeps_counting_from_posted_on() -> None:
    job = {"posted_on": "2026-09-14", "posted_date": "3 days ago"}
    assert posted_age_days(job, today=date(2026, 9, 24)) == 10
    assert not is_within_freshness(job, 7, today=date(2026, 9, 24))


def test_posted_age_handles_naukri_variants() -> None:
    assert posted_age_days({"posted_date": "3+ weeks ago"}) == 21
    assert posted_age_days({"posted_date": "30+ days ago"}) == 30
    assert posted_age_days({"posted_date": "Just now"}) == 0
    assert posted_age_days({"posted_date": "Few hours ago"}) == 0
    assert posted_age_days({"posted_date": "1 month ago"}) == 30
