"""
End-to-end test of the pipeline stages with a fake discoverer, so no
browser or network is needed.
"""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.config import load_matching_rules, load_settings
from src.database import Database
from src.models import Job
from src.pipelines.discovery_pipeline import RunContext, run
from src.resume_profile import load_resume_profile

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 9, 24)

GOOD_DESCRIPTION = (
    "Gather business requirements, write BRDs and user stories, coordinate "
    "UAT, work with stakeholders, and use SQL for product analytics in an "
    "AI-powered enterprise SaaS platform."
)


def make_jobs() -> list[Job]:
    return [
        Job(
            title="AI Business Analyst",
            company="Fresh Fit",
            location="Hyderabad",
            job_url="https://example.com/jobs/fresh-fit",
            source="naukri",
            experience="3-5 Yrs",
            posted_date="2 days ago",
            description=GOOD_DESCRIPTION,
        ),
        Job(
            title="Business Analyst",
            company="Stale Co",
            location="Hyderabad",
            job_url="https://example.com/jobs/stale",
            source="naukri",
            experience="3-5 Yrs",
            posted_date="3 weeks ago",
            description=GOOD_DESCRIPTION,
        ),
        Job(
            title="Lead Business Analyst",
            company="Senior Co",
            location="Hyderabad",
            job_url="https://example.com/jobs/too-senior",
            source="naukri",
            experience="8-13 Yrs",
            posted_date="1 day ago",
            description=GOOD_DESCRIPTION,
        ),
        Job(
            title="Java Developer",
            company="Other Co",
            location="Hyderabad",
            job_url="https://example.com/jobs/unrelated",
            source="naukri",
            posted_date="1 day ago",
            description="Spring Boot microservices and Kafka.",
        ),
    ]


@pytest.fixture
def context(tmp_path) -> RunContext:
    settings = load_settings(PROJECT_ROOT / "config" / "settings.yaml")
    settings.project.database_path = str(tmp_path / "jobs.db")
    settings.project.export_directory = str(tmp_path / "exports")
    settings.project.log_directory = str(tmp_path / "logs")
    settings.application.prepare_after_search = False
    return RunContext(
        project_root=PROJECT_ROOT,
        settings=settings,
        rules=load_matching_rules(PROJECT_ROOT / "config" / "matching_rules.yaml"),
        profile=load_resume_profile(PROJECT_ROOT / "config" / "candidate_profile.yaml"),
    )


def fake_discover(jobs):
    def discover(context, known_urls, stats):
        found = [job for job in jobs if job.job_url not in known_urls]
        stats.discovered = len(found)
        return found

    return discover


def test_pipeline_filters_saves_and_exports(context):
    stats = run(context, discover=fake_discover(make_jobs()), today=TODAY)

    assert stats.discovered == 4
    assert stats.stale == 1
    assert stats.rejected == 2
    assert stats.new == 1

    with Database(context.database_path) as database:
        saved = database.get_jobs()
    assert [job["company"] for job in saved] == ["Fresh Fit"]
    assert saved[0]["experience_min"] == 3
    assert saved[0]["posted_on"] == "2026-09-22"

    daily = load_workbook(context.export_directory / "new_leads_7d_2026-09-24.xlsx").active
    header = [cell.value for cell in daily[1]]
    assert daily.max_row == 2
    assert daily[2][header.index("company")].value == "Fresh Fit"
    assert daily[2][header.index("experience")].value == "3-5 Yrs"

    assert (context.export_directory / "master_jobs.xlsx").exists()
    assert (context.export_directory / "application_review_2026-09-24.html").exists()


def test_second_run_exports_no_repeat_leads(context):
    run(context, discover=fake_discover(make_jobs()), today=TODAY)
    stats = run(context, discover=fake_discover(make_jobs()), today=TODAY)

    assert stats.new == 0
    daily = load_workbook(context.export_directory / "new_leads_7d_2026-09-24.xlsx").active
    assert daily.max_row == 1
