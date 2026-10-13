"""Tailored resume / cover letter generation, with a fake LLM (no network)."""

from __future__ import annotations

import copy
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from src.config import TailoringSettings, load_matching_rules, load_settings
from src.database import Database
from src.models import Job
from src.pipelines.discovery_pipeline import RunContext, run
from src.resume_profile import load_resume_profile
from src.tailoring.generator import DocumentGenerator, GenerationResult
from src.tailoring.llm import GeminiClient, LLMError, parse_json_text
from src.tailoring.master_resume import MasterResumeError, load_master_resume
from src.tailoring.tailor import TailoredApplication, fact_check, tailor_application

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 9, 24)

MASTER_RESUME = """\
Asha Rao
asha@example.com | +91 90000 00000 | Hyderabad

SUMMARY
Business analyst with experience in AI-powered analytics and enterprise SaaS.

SKILLS
Business Analysis, Requirements Gathering, SQL, AWS Athena, Stakeholder Management,
UAT, User Stories, Product Analytics, Jira

EXPERIENCE
Business Analyst - AI Analytics | Nimbus Analytics | Hyderabad | Jan 2024 - Present
- Gathered business requirements and wrote BRDs and user stories for an AI analytics platform.
- Coordinated UAT with 12 enterprise clients.
- Queried AWS Athena with SQL to build product usage dashboards.

Customer Experience Analyst | Concentrix | Hyderabad | Jun 2021 - Dec 2023
- Analysed customer feedback for a credit card portfolio and reported trends to stakeholders.

EDUCATION
B.Com | Osmania University | 2017
"""

GOOD_DESCRIPTION = (
    "We need an AI Business Analyst to gather business requirements, write user "
    "stories and BRDs, coordinate UAT and work with stakeholders. Strong SQL and "
    "product analytics skills needed; experience with Power BI and Python preferred. "
    "You will partner with engineering on an AI-powered enterprise SaaS platform. "
) * 2


def model_answer(**overrides):
    answer = {
        "candidate": {"name": "Asha Rao", "contact_line": "asha@example.com | Hyderabad"},
        "headline": "AI Business Analyst",
        "summary": "Business analyst with AI-powered analytics experience. Increased revenue by 45%.",
        "skills": ["Business Analysis", "SQL", "Requirements Gathering", "Power BI", "UAT"],
        "experience": [
            {
                "title": "Business Analyst - AI Analytics",
                "company": "Nimbus Analytics",
                "location": "Hyderabad",
                "dates": "Jan 2024 - Present",
                "bullets": [
                    "Gathered business requirements and wrote BRDs and user stories.",
                    "Coordinated UAT with 12 enterprise clients.",
                    "Cut reporting time by 70% using SQL.",
                ],
            },
            {
                "title": "Senior Analyst",
                "company": "Invented Corp",
                "dates": "2019 - 2021",
                "bullets": ["Did things."],
            },
        ],
        "projects": [],
        "education": [{"degree": "B.Com", "institution": "Osmania University", "dates": "2017"}],
        "certifications": ["PMP"],
        "cover_letter": {
            "salutation": "Dear Hiring Manager,",
            "paragraphs": [
                "I am applying for the AI Business Analyst role.",
                "At Nimbus Analytics I coordinated UAT with 12 enterprise clients.",
                "I would welcome a conversation.",
            ],
            "closing": "Sincerely,",
        },
        "jd_keywords": ["business requirements", "user stories", "UAT", "SQL", "Power BI", "Python"],
        "missing_keywords": ["Power BI", "Python"],
    }
    answer.update(overrides)
    return answer


class FakeClient:
    model_name = "fake-model"

    def __init__(self, answer=None, error: Exception | None = None):
        self.answer = answer or model_answer()
        self.error = error
        self.calls = 0

    def generate_json(self, system, prompt):
        self.calls += 1
        if self.error:
            raise self.error
        assert "MASTER RESUME" in prompt
        return copy.deepcopy(self.answer)


def settings(tmp_path: Path, **overrides) -> TailoringSettings:
    values = dict(
        enabled=True,
        output_directory=str(tmp_path / "applications"),
        seconds_between_requests=0,
        min_description_chars=100,
    )
    values.update(overrides)
    return TailoringSettings(**values)


def make_job(url: str, score: float, **extra) -> dict:
    job = {
        "job_url": url,
        "title": "AI Business Analyst",
        "company": "Acme Labs",
        "location": "Hyderabad",
        "match_score": score,
        "match_category": "good",
        "description": GOOD_DESCRIPTION,
        "status": "new",
    }
    job.update(extra)
    return job


def generator(tmp_path, client=None, **overrides) -> DocumentGenerator:
    return DocumentGenerator(
        settings(tmp_path, **overrides),
        tmp_path,
        client or FakeClient(),
        master_resume=MASTER_RESUME,
        contact={},
        sleep=lambda seconds: None,
        today=TODAY,
    )


# ----------------------------------------------------------------------
# Fact check
# ----------------------------------------------------------------------


def test_fact_check_removes_invented_content():
    application = fact_check(TailoredApplication.model_validate(model_answer()), MASTER_RESUME)

    assert "Power BI" not in application.skills
    assert {"Business Analysis", "SQL", "UAT"} <= set(application.skills)
    assert [entry.company for entry in application.experience] == ["Nimbus Analytics"]
    bullets = application.experience[0].bullets
    assert "Coordinated UAT with 12 enterprise clients." in bullets
    assert not any("70%" in bullet for bullet in bullets)
    assert application.certifications == []
    assert "45" not in application.summary
    assert application.education[0].institution == "Osmania University"
    assert len(application.fact_check_warnings) >= 4


def test_profile_evidence_allows_listed_skill():
    application = fact_check(
        TailoredApplication.model_validate(model_answer()),
        MASTER_RESUME,
        extra_evidence=["Power BI"],
    )
    assert "Power BI" in application.skills


def test_tailor_application_scores_keywords():
    application = tailor_application(make_job("u", 70), MASTER_RESUME, FakeClient())
    assert "Power BI" in application.unmatched_keywords
    assert "SQL" in application.matched_keywords
    assert 0 < application.keyword_coverage < 100


def test_tailor_application_rejects_empty_experience():
    answer = model_answer(experience=[{"title": "X", "company": "Invented Corp", "bullets": ["a"]}])
    with pytest.raises(LLMError):
        tailor_application(make_job("u", 70), MASTER_RESUME, FakeClient(answer))


# ----------------------------------------------------------------------
# Selection and generation
# ----------------------------------------------------------------------


def test_only_jobs_at_or_above_threshold_are_selected(tmp_path):
    jobs = [
        make_job("a", 59.9),
        make_job("b", 60),
        make_job("c", 85),
        make_job("d", 90, status="applied"),
        make_job("e", 80, description="short"),
        make_job("f", 75),
        make_job("g", None),
    ]
    result = GenerationResult()
    selected = generator(tmp_path, max_jobs_per_run=2).select_jobs(jobs, {"f"}, result)

    assert [job["job_url"] for job in selected] == ["c", "b"]
    assert result.below_threshold == 2
    assert result.already_done == 1
    assert result.no_description == 1
    assert result.deferred == 0


def test_limit_defers_lowest_scores(tmp_path):
    jobs = [make_job(str(i), 60 + i) for i in range(5)]
    result = GenerationResult()
    selected = generator(tmp_path, max_jobs_per_run=3).select_jobs(jobs, set(), result)
    assert [job["job_url"] for job in selected] == ["4", "3", "2"]
    assert result.deferred == 2


def test_generate_writes_docx_and_skips_on_rerun(tmp_path):
    client = FakeClient()
    with Database(tmp_path / "jobs.db") as database:
        result = generator(tmp_path, client).generate(database, [make_job("https://x/1", 72)])
        assert len(result.created) == 1
        record = database.get_tailored_documents()["https://x/1"]

        resume = Document(record["resume_path"])
        text = "\n".join(p.text for p in resume.paragraphs)
        assert "Asha Rao" in text
        assert "PROFESSIONAL EXPERIENCE" in text
        assert "Invented Corp" not in text
        assert not resume.tables  # ATS: no tables

        letter = "\n".join(p.text for p in Document(record["cover_letter_path"]).paragraphs)
        assert "Acme Labs" in letter
        assert "Re: Application for AI Business Analyst" in letter

        report = Path(record["report_path"]).read_text(encoding="utf-8")
        assert "Power BI" in report

        again = generator(tmp_path, client).generate(database, [make_job("https://x/1", 72)])
        assert again.already_done == 1
        assert client.calls == 1

        # Deleting the files makes the job eligible again.
        Path(record["resume_path"]).unlink()
        generator(tmp_path, client).generate(database, [make_job("https://x/1", 72)])
        assert client.calls == 2


def test_contact_from_application_profile_wins(tmp_path):
    gen = generator(tmp_path)
    gen.contact = {"full_name": "Asha R", "email": "a@b.com", "phone": "123"}
    with Database(tmp_path / "jobs.db") as database:
        gen.generate(database, [make_job("u", 70)])
        record = database.get_tailored_documents()["u"]
    assert "Asha_R_Resume" in Path(record["resume_path"]).name
    text = "\n".join(p.text for p in Document(record["resume_path"]).paragraphs)
    assert "a@b.com | 123" in text


def test_failures_are_not_recorded_and_stop_after_three(tmp_path):
    client = FakeClient(error=LLMError("HTTP 429: quota"))
    jobs = [make_job(str(i), 70) for i in range(5)]
    with Database(tmp_path / "jobs.db") as database:
        result = generator(tmp_path, client).generate(database, jobs)
        assert database.get_tailored_documents() == {}
    assert len(result.failed) == 3
    assert result.deferred == 2
    assert client.calls == 3


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------


def test_parse_json_text_handles_code_fences():
    assert parse_json_text('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json_text('Here: {"a": 2} done') == {"a": 2}
    with pytest.raises(LLMError):
        parse_json_text("no json")


def test_gemini_extracts_text_and_requires_key():
    data = {"candidates": [{"content": {"parts": [{"text": '{"x"'}, {"text": ": 1}"}]}}]}
    assert GeminiClient._extract_text(data) == '{"x": 1}'
    with pytest.raises(LLMError):
        GeminiClient._extract_text({"promptFeedback": {"blockReason": "SAFETY"}})
    with pytest.raises(LLMError):
        GeminiClient(api_key="", model="m")


def test_load_master_resume_docx_and_missing(tmp_path):
    path = tmp_path / "resume.docx"
    document = Document()
    for line in MASTER_RESUME.splitlines():
        document.add_paragraph(line)
    document.save(str(path))
    assert "Nimbus Analytics" in load_master_resume(path)

    with pytest.raises(MasterResumeError):
        load_master_resume(tmp_path / "missing.docx")


# ----------------------------------------------------------------------
# Pipeline + Excel
# ----------------------------------------------------------------------


def test_pipeline_links_documents_in_excel(tmp_path):
    resume_path = tmp_path / "master.txt"
    resume_path.write_text(MASTER_RESUME, encoding="utf-8")

    app_settings = load_settings(PROJECT_ROOT / "config" / "settings.yaml")
    app_settings.project.database_path = str(tmp_path / "jobs.db")
    app_settings.project.export_directory = str(tmp_path / "exports")
    app_settings.project.log_directory = str(tmp_path / "logs")
    app_settings.application.prepare_after_search = False
    app_settings.tailoring = settings(tmp_path, master_resume_path=str(resume_path), min_match_score=0)
    context = RunContext(
        project_root=tmp_path,
        settings=app_settings,
        rules=load_matching_rules(PROJECT_ROOT / "config" / "matching_rules.yaml"),
        profile=load_resume_profile(PROJECT_ROOT / "config" / "candidate_profile.yaml"),
    )
    job = Job(
        title="AI Business Analyst",
        company="Fresh Fit",
        location="Hyderabad",
        job_url="https://example.com/jobs/fresh-fit",
        source="naukri",
        experience="3-5 Yrs",
        posted_date="2 days ago",
        description=GOOD_DESCRIPTION,
    )

    run(context, discover=lambda ctx, known, stats: [job], today=TODAY, llm_client=FakeClient())

    for name in ("new_leads_7d_2026-09-24.xlsx", "master_jobs.xlsx"):
        sheet = load_workbook(tmp_path / "exports" / name).active
        headers = [cell.value for cell in sheet[1]]
        cell = sheet.cell(row=2, column=headers.index("tailored_resume") + 1)
        assert cell.value.endswith(".docx")
        assert cell.hyperlink.target.startswith("file:")
        letter = sheet.cell(row=2, column=headers.index("cover_letter") + 1)
        assert "Cover_Letter" in letter.value
