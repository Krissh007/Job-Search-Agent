from __future__ import annotations

import json
import re
from dataclasses import asdict, is_dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


class JobStorage:
    """
    Saves each discovered job as a JSON file and exports
    all jobs to an Excel workbook.

    No database is used.
    """

    def __init__(self, project_root: Path, config: dict):
        self.project_root = Path(project_root)

        project_config = config.get("project", {})

        self.data_directory = self.project_root / "data"

        self.jobs_directory = self.data_directory / "jobs"

        self.export_directory = (
            self.project_root
            / project_config.get(
                "export_directory",
                "data/exports",
            )
        )

        self.jobs_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        self.export_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

    @staticmethod
    def _safe_filename(value: str) -> str:
        """
        Converts text into a Windows-safe filename.
        """

        value = value.strip()

        value = re.sub(
            r'[<>:"/\\|?*]',
            "_",
            value,
        )

        value = re.sub(
            r"\s+",
            "_",
            value,
        )

        return value[:120] or "job"

    @staticmethod
    def _job_to_dict(job: Any) -> dict:
        """
        Converts a dataclass, dictionary, or normal object
        into a dictionary.
        """

        if isinstance(job, dict):
            return dict(job)

        if is_dataclass(job):
            return asdict(job)

        result = {}

        for field in [
            "title",
            "company",
            "location",
            "job_url",
            "source",
            "description",
            "posted_date",
            "salary",
            "job_type",
            "match_score",
            "match_level",
            "matching_roles",
            "matching_skills",
            "missing_skills",
            "matched_responsibilities",
            "matched_tools",
            "matched_industries",
            "preferred_location_match",
            "match_reason",
        ]:
            result[field] = getattr(job, field, "")

        return result

    @staticmethod
    def _job_key(job_data: dict) -> str:
        """
        Creates a stable key for a job.

        The job URL is preferred because it is normally unique.
        """

        job_url = str(
            job_data.get("job_url", "")
        ).strip()

        if job_url:
            return job_url

        return "|".join(
            [
                str(job_data.get("title", "")).strip().lower(),
                str(job_data.get("company", "")).strip().lower(),
                str(job_data.get("location", "")).strip().lower(),
            ]
        )

    @staticmethod
    def _excel_value(value: Any) -> Any:
        """
        Converts lists and dictionaries into readable Excel text.
        """

        if isinstance(value, list):
            return ", ".join(str(item) for item in value)

        if isinstance(value, dict):
            return json.dumps(
                value,
                ensure_ascii=False,
            )

        if isinstance(value, bool):
            return "Yes" if value else "No"

        return value

    def save_jobs(self, jobs: Iterable) -> int:
        """
        Saves jobs as individual JSON files.

        Existing jobs are not overwritten if the same job URL
        or fallback identity already exists.

        Returns the number of newly saved jobs.
        """

        saved_count = 0
        existing_keys = set()

        for json_file in self.jobs_directory.glob("*.json"):
            try:
                with json_file.open(
                    "r",
                    encoding="utf-8",
                ) as file:
                    existing_job = json.load(file)

                existing_keys.add(
                    self._job_key(existing_job)
                )

            except (OSError, json.JSONDecodeError):
                continue

        for job in jobs:
            job_data = self._job_to_dict(job)

            job_key = self._job_key(job_data)

            if not job_key or job_key in existing_keys:
                continue

            timestamp = datetime.now().strftime(
                "%Y%m%d_%H%M%S_%f"
            )

            title = self._safe_filename(
                str(job_data.get("title", "job"))
            )

            company = self._safe_filename(
                str(
                    job_data.get(
                        "company",
                        "unknown_company",
                    )
                )
            )

            filename = (
                f"{timestamp}_{company}_{title}.json"
            )

            output_path = self.jobs_directory / filename

            try:
                with output_path.open(
                    "w",
                    encoding="utf-8",
                ) as file:
                    json.dump(
                        job_data,
                        file,
                        indent=2,
                        ensure_ascii=False,
                    )

                existing_keys.add(job_key)
                saved_count += 1

            except OSError as error:
                print(
                    f"Could not save job file "
                    f"{output_path}: {error}"
                )

        return saved_count

    def load_all_jobs(self) -> list[dict]:
        """
        Loads all saved JSON job files.
        """

        jobs = []

        for json_file in sorted(
            self.jobs_directory.glob("*.json")
        ):
            try:
                with json_file.open(
                    "r",
                    encoding="utf-8",
                ) as file:
                    job_data = json.load(file)

                jobs.append(job_data)

            except (OSError, json.JSONDecodeError) as error:
                print(
                    f"Could not read {json_file}: {error}"
                )

        return jobs

    def export_to_excel(
        self,
        output_filename: str = "job_search_results.xlsx",
    ) -> Path:
        """
        Exports all JSON job files to Excel.
        """

        jobs = self.load_all_jobs()

        output_path = (
            self.export_directory / output_filename
        )

        columns = [
            "title",
            "company",
            "location",
            "match_score",
            "match_level",
            "matching_roles",
            "matching_skills",
            "missing_skills",
            "matched_responsibilities",
            "matched_tools",
            "matched_industries",
            "preferred_location_match",
            "match_reason",
            "job_url",
            "source",
            "description",
            "posted_date",
            "salary",
            "job_type",
        ]

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Jobs"

        worksheet.append(columns)

        for job in jobs:
            worksheet.append(
                [
                    self._excel_value(
                        job.get(column, "")
                    )
                    for column in columns
                ]
            )

        # Format the header.
        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        worksheet.freeze_panes = "A2"

        if jobs:
            worksheet.auto_filter.ref = worksheet.dimensions

        # Set reasonable column widths.
        for column_index, column_name in enumerate(
            columns,
            start=1,
        ):
            column_letter = get_column_letter(
                column_index
            )

            max_length = len(column_name)

            for row in worksheet.iter_rows(
                min_col=column_index,
                max_col=column_index,
            ):
                for cell in row:
                    if cell.value is not None:
                        max_length = max(
                            max_length,
                            len(str(cell.value)),
                        )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 60)

        # Make the match score column easier to read.
        match_score_column = columns.index(
            "match_score"
        ) + 1

        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(
                row=row,
                column=match_score_column,
            )

            if isinstance(cell.value, (int, float)):
                cell.number_format = "0"

        workbook.save(output_path)

        return output_path

    def get_job_count(self) -> int:
        """
        Returns the number of saved JSON job files.
        """

        return len(
            list(
                self.jobs_directory.glob("*.json")
            )
        )