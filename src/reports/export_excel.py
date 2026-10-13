from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.freshness import is_within_freshness, posted_age_days
from src.database import Database


EXPORT_COLUMNS = [
    "title",
    "company",
    "location",
    "experience",
    "salary",
    "source",
    "posted_date",
    "posted_on",
    "freshness_age_days",
    "match_score",
    "match_category",
    "match_reasons",
    "missing_keywords",
    "search_term",
    "job_url",
    "tailored_resume",
    "cover_letter",
    "application_progress",
    "description",
    "discovered_at",
]


# Columns rendered as clickable links. Document columns show the file name
# and link to the file itself.
LINK_COLUMNS = ("job_url", "tailored_resume", "cover_letter", "application_progress")
FILE_LINK_COLUMNS = ("tailored_resume", "cover_letter")


def attach_tailored_documents(
    row: dict[str, Any],
    documents: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Copy tailored resume / cover-letter paths for this job onto the row."""
    record = documents.get(str(row.get("job_url") or ""))
    if record:
        row["tailored_resume"] = record["resume_path"]
        row["cover_letter"] = record["cover_letter_path"]
    return row


def _write_links(worksheet: Any, columns: list[str], row: dict[str, Any]) -> None:
    for name in LINK_COLUMNS:
        target = row.get(name)
        if not target or name not in columns:
            continue
        cell = worksheet.cell(row=worksheet.max_row, column=columns.index(name) + 1)
        target = str(target)
        if name in FILE_LINK_COLUMNS:
            cell.value = Path(target).name
            target = Path(target).resolve().as_uri()
        cell.hyperlink = target
        cell.style = "Hyperlink"


def _cell_value(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(
            f"{key}: {item}" for key, item in value.items()
        )
    return "" if value is None else value


def _save_workbook(workbook: Workbook, path: Path) -> Path:
    """Save, or write a timestamped copy when the file is open in Excel."""
    try:
        workbook.save(path)
        return path
    except PermissionError:
        fallback_path = path.with_name(
            f"{path.stem}_updated_{datetime.now().strftime('%H%M%S')}{path.suffix}"
        )
        workbook.save(fallback_path)
        print(f"{path.name} is locked; wrote fallback file: {fallback_path}")
        return fallback_path


def export_daily_new_leads(
    jobs: Iterable[dict[str, Any]],
    export_directory: str | Path,
    run_date: date | None = None,
    freshness_windows: Iterable[int] = (7, 30),
) -> list[Path]:
    """Write one dated Excel workbook per freshness window."""
    export_path = Path(export_directory)
    export_path.mkdir(parents=True, exist_ok=True)
    current_date = run_date or date.today()
    job_list = list(jobs)
    output_paths: list[Path] = []

    for window_days in freshness_windows:
        window = int(window_days)
        path = export_path / (
            f"new_leads_{window}d_{current_date.isoformat()}.xlsx"
        )
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = f"New leads {window}d"
        worksheet.append(EXPORT_COLUMNS)

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        exported_count = 0
        for job in job_list:
            if not is_within_freshness(job, window, current_date):
                continue

            row = dict(job)
            row["freshness_age_days"] = posted_age_days(job, current_date)
            worksheet.append(
                [
                    _cell_value(row.get(column))
                    for column in EXPORT_COLUMNS
                ]
            )
            _write_links(worksheet, EXPORT_COLUMNS, row)
            exported_count += 1

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for column_index, column_name in enumerate(EXPORT_COLUMNS, start=1):
            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = min(max(len(column_name) + 2, 14), 32)

        path = _save_workbook(workbook, path)
        output_paths.append(path)
        print(
            f"Excel export completed: {path} "
            f"({exported_count} new leads)"
        )

    return output_paths


def export_master_jobs(
    database_path: str | Path,
    export_directory: str | Path,
) -> Path:
    """Export one deduplicated workbook from the SQLite source of truth."""
    output_directory = Path(export_directory)
    output_directory.mkdir(parents=True, exist_ok=True)
    output_path = output_directory / "master_jobs.xlsx"

    with Database(database_path) as database:
        jobs = database.get_jobs()
        application_by_url = database.get_applications()
        documents = database.get_tailored_documents()

    unique_jobs: dict[str, dict[str, Any]] = {}
    for job in jobs:
        if job.get("status") == "rejected":
            continue
        job_url = str(job.get("job_url") or "")
        if not job_url:
            continue
        row = dict(job)
        application = application_by_url.get(job_url)
        if application:
            row["application_progress"] = application["review_url"]
            row["application_status"] = application["status"]
        attach_tailored_documents(row, documents)
        unique_jobs[job_url] = row

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Master jobs"
    columns = EXPORT_COLUMNS + ["application_status"]
    worksheet.append(columns)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in unique_jobs.values():
        row["freshness_age_days"] = posted_age_days(row)
        worksheet.append([_cell_value(row.get(column)) for column in columns])
        _write_links(worksheet, columns, row)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    output_path = _save_workbook(workbook, output_path)
    print(f"Master Excel export completed: {output_path} ({len(unique_jobs)} jobs)")
    return output_path
